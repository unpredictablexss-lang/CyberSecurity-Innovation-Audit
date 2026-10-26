from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent
IMG = ROOT.parent / "deliveries" / "chapter6_20261022" / "images"
SOURCE = Path(r"C:\Users\DeepSleep\xwechat_files\wxid_txs8nygg8zhb22_5718\msg\file\2026-10\中国研究生网络安全创新大赛作品报告(1).docx")
OUT = ROOT.parent / "deliveries" / "chapter6_20261022" / "第六章_系统功能实现与用户交互.docx"

doc = Document(SOURCE)
body = doc._element.body
sectPr = deepcopy(body.sectPr)
for child in list(body):
    body.remove(child)
body.append(sectPr)

# Keep the template's A4 page geometry and its existing style/theme package.
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin, sec.bottom_margin = Inches(1), Inches(1)
sec.left_margin, sec.right_margin = Inches(1.12), Inches(1.02)
sec.header_distance, sec.footer_distance = Inches(.45), Inches(.45)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(10.5)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.first_line_indent = Inches(.25)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

h1 = doc.styles["Heading 1"]
h1.font.name = "黑体"
h1.font.size = Pt(22)
h1.font.bold = True
h1.font.color.rgb = None
h1._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
h1.paragraph_format.space_before = Pt(0)
h1.paragraph_format.space_after = Pt(16)
h1.paragraph_format.keep_with_next = True

for sty_name, size in [("Heading 2", 14), ("Heading 3", 12)]:
    try:
        st = doc.styles[sty_name]
    except KeyError:
        st = doc.styles.add_style(sty_name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = normal
    st.font.name = "黑体"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = None
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    st.paragraph_format.space_before = Pt(8)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.keep_with_next = True

try:
    caption = doc.styles["Caption"]
except KeyError:
    caption = doc.styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption.base_style = normal
caption.font.name = "宋体"
caption.font.size = Pt(10.5)
caption._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
caption.paragraph_format.space_before = Pt(3)
caption.paragraph_format.space_after = Pt(8)
caption.paragraph_format.keep_with_next = False

# Replace template's stale static page label with current page fields.
footer = sec.footer
for p in footer.paragraphs:
    p.clear()
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.first_line_indent = Inches(0)
fp.add_run("第 ")
for field in ("PAGE",):
    r = fp.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = field
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t"); txt.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    r._r.extend([begin, instr, sep, txt, end])
fp.add_run(" 页  共 ")
r = fp.add_run()
for tag, value in [("begin", None), (None, "NUMPAGES"), ("separate", None), ("text", "8"), ("end", None)]:
    el = OxmlElement("w:fldChar" if tag in ("begin", "separate", "end") else "w:instrText" if tag is None else "w:t")
    if tag in ("begin", "separate", "end"): el.set(qn("w:fldCharType"), tag)
    elif tag is None: el.set(qn("xml:space"), "preserve"); el.text = value
    else: el.text = value
    r._r.append(el)
fp.add_run(" 页")

def para(text="", style="Normal", bold_lead=None):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p

def heading(text):
    return doc.add_paragraph(text, style="Heading 2")

def figure(filename, caption_text, width=6.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(str(IMG / filename), width=Inches(width))
    c = doc.add_paragraph(caption_text, style="Caption")
    c.paragraph_format.first_line_indent = Inches(0)
    return c

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr=t._tbl.tblPr
    borders=OxmlElement("w:tblBorders")
    for edge in ("top","left","bottom","right","insideH","insideV"):
        el=OxmlElement("w:"+edge); el.set(qn("w:val"),"single"); el.set(qn("w:sz"),"4"); el.set(qn("w:color"),"A6A6A6"); borders.append(el)
    tblPr.append(borders)
    for i, text in enumerate(headers):
        cell=t.rows[0].cells[i]; cell.text=text; cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs: run.bold=True
        shade=OxmlElement("w:shd"); shade.set(qn("w:fill"), "D9EAF7"); cell._tc.get_or_add_tcPr().append(shade)
    for row in rows:
        cells=t.add_row().cells
        for i,text in enumerate(row):
            cells[i].text=text; cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in t.rows:
        for i,cell in enumerate(row.cells):
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent=Inches(0)
                p.paragraph_format.line_spacing=1.15
                p.paragraph_format.space_after=Pt(0)
                for run in p.runs:
                    run.font.name="Times New Roman"; run.font.size=Pt(9.5); run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"宋体")
    return t

doc.add_paragraph("第六章 系统功能实现与用户交互", style="Heading 1")
para("本章介绍 OHScanner 的源码接入、审计任务管理、风险研判与报告导出流程。演示基于 OpenHarmony libpng 已保存的静态扫描任务，仓库为 openharmony/third_party_libpng，扫描提交为 e57b1079564a02a50321ec528aa6e41884979d32。章节结合实际运行界面说明系统功能，并以任务记录中的候选、覆盖及复核状态作为结果依据。")

heading("6.1 用户视角的一体化工作台")
para("工作台将仓库登记、审计创建、任务跟踪、风险筛选、人工调查、PoC 验证和报告导出整合于统一导航。总览集中呈现本地仓库、已保存扫描及待复核风险，并提供服务能力入口。用户既可从总览创建审计，也可从风险候选回溯固定源码和证据记录。任务状态“完成（有限制）”表示扫描已生成结果，同时报告列明解析或覆盖范围。")
figure("overview.png", "图6-1  总览页展示工作台入口、仓库与任务摘要（更新前端，2026-10-22）")

heading("6.2 仓库目录、源码同步与版本选择")
para("系统提供在线仓库目录和本地仓库两类源码入口。在线目录支持检索公开仓库元数据、查看默认分支及预览完整提交；用户确认下载后，可选择同步完成即固定版本并排队静态扫描，实现源码获取与任务创建的连续操作。当前演示展示目录元数据及同步配置界面；本次扫描对象来自已登记的本地仓库。")
para("本地接入通过受控源码根目录或登记本地 Git 仓库完成。路径需位于管理员登记的根目录中；系统检查 Git 可读性、仓库状态和 HEAD，失败或越界时不自动放宽目录范围。扫描任务保存完整 Commit、Git Tree、纳入路径及报告哈希。分支名只是选择入口，任务实际绑定到解析后的提交；之后分支移动不会改变已保存任务的源码身份。")
figure("repositories-local.png", "图6-2  已登记仓库与源码同步入口；路径和仓库身份均需按界面提示核验")

heading("6.3 三种审计范围与检测器选择")
para("审计范围定义源码版本的选取与比较方式，研判流程定义候选的分析方式。仓库快照针对单一固定提交开展基线审查；时间增量依据提交者时间窗口确定并固定两个端点，用于分析区间变化；跨版本安全分析比较同一仓库的两个分支、Tag 或 Commit。两类比较任务均保存端点版本，为候选变化分析提供可复现的源码基准。")
para("创建页面在确定仓库和版本后，支持配置审计方式、纳入路径、解析选项及规则范围。静态模式运行本地检测器；智能审计在展示候选发送范围后取得本次授权。任务报告分别记录规则计划、实际运行状态和文件解析覆盖；系统能力页则列出已接入检测器及其适用范围。对未启用规则、解析不完整或未运行任务，报告采用 unknown/partial 状态标注。")
table(["审计范围", "固定对象", "主要用途", "结果解释"], [
    ["仓库快照", "一个 Commit / Tree", "基线审查与版本留档", "本次版本候选，不代表漏洞确认"],
    ["时间增量", "时间窗口内的两个端点", "查看区间变化", "不是逐提交扫描；需复核端点差异"],
    ["跨版本", "同仓库的两个分支、Tag 或 Commit", "版本或分支比较", "候选变化不自动归因为引入或修复"],
])
figure("audit-repository.png", "图6-3  仓库快照审计创建页：固定仓库后配置本次扫描")
figure("audit-time.png", "图6-4  时间增量页：按时间窗口预览并固定两个端点")
figure("audit-versions.png", "图6-5  跨版本页：选择基准版本与目标版本")

heading("6.4 审计创建、任务进度和结果查看")
para("典型操作顺序为：选择仓库，解析并固定版本，选择快照、时间或跨版本范围，核对扫描路径与审计方式，再提交任务。任务中心汇总源码扫描和区间操作。任务详情记录整体状态、Commit、Tree、源码文件数量、耗时、规则计划、报告哈希以及各处理阶段。用户遇到失败或有限制状态时，可先查看阶段和覆盖说明，再决定重试或缩小范围；重试不会改变原任务固定的版本记录。")
para("libpng 固定扫描于 2026-10-22 完成，任务状态为“静态审计完成（有限制）”，耗时 36.49 秒。任务纳入 99 个源码文件，其中 38 个完整解析、61 个解析不完整；检测器产生 68 项原始候选观测，经报告归并为 29 组，任务记录的人工确认漏洞数为 0。该统计呈现了本次任务的处理规模、解析覆盖和候选归并结果。")
figure("tasks.png", "图6-6  任务中心列出扫描任务、状态与查看入口")
figure("task-detail-current.png", "图6-7  libpng 固定扫描详情：提交、Tree、覆盖限制和报告摘要")

heading("6.5 风险中心、候选证据与风险分层")
para("风险中心支持按可信程度、初判危害、漏洞状态、仓库与 Commit、检测器及复核状态筛选候选。可信程度用于安排调查优先级，危害等级描述问题成立时的潜在影响，复核状态记录人工调查进度，三类信息分别呈现，支持研究人员快速定位待处理事项。")
para("libpng 任务在 arm/palette_neon_intrinsics.c:63 与 :108 两处生成 CVE-2026-33636 候选，分别对应同一问题的源码位置。GitHub 上游公告列出的影响版本为 1.6.36–1.6.55，受影响范围为启用硬件优化的 ARM/AArch64 Neon 调色板展开代码；通用 C 路径及若干其他 SIMD 路径不受影响。本次源码版本为 libpng 1.6.55，与公告版本范围相符。产品构建是否包含该组件、是否启用 Neon 及相关输入路径是否可达，仍需结合目标构建配置核验。[1]")
figure("libpng-risk-evidence-level.png", "图6-8  风险候选按证据状态与初判危害分层；层级不等于漏洞确认")
figure("candidate-evidence-current.png", "图6-9  CVE-2026-33636 候选详情：同一固定提交上的源码位置及待核实义务")

heading("6.6 三种研判流程与授权交互")
para("研判流程与审计范围相互独立。静态扫描加人工审计只在本地运行检测器，系统提供风险候选和源码证据，研究人员自行复核，不调用大模型。先本地排序、再批量 AI 审计时，系统在本地归并或排序候选；用户选择需要调查的范围，查看发送内容、端点和预算后，才能授权快筛或进一步深查。一键智能审计则在创建任务时先取得授权，静态扫描结束后自动发起初步模型研判，并按策略对符合条件的候选继续工具型追查。")
para("渐进流程包括 L0/L1 本地整理与排序、L2 初步模型研判和 L3 只读工具追查。各阶段分别记录计划、排队、完成、失败或未调度状态；用户可在授权前核对候选、源码片段、目标端点及范围上限，授权范围绑定本次固定任务。模型输出作为辅助研判信息，与检测证据及人工复核记录共同呈现。当前演示环境未配置模型服务，因此本例 AI 阶段状态为 not_run，源码未发送至模型端点。")

heading("6.7 人工复核、案例管理与 PoC 验证")
para("候选详情提供固定源码摘录、Commit、源码哈希、检测器版本、规则标识及扫描快照哈希。人工复核可围绕组件身份、构建配置、攻击者可控输入路径和影响证据开展，并将结果记录为确认、误报、调查中或证据不足；复核说明关联文件位置、判断依据与后续事项，审查记录按时间追加。当前 libpng 候选处于“待复核”状态，产品构建、输入可达性及影响证据尚待核实。")
para("调查案例用于组织相关候选、复核意见与验证过程。PoC 验证页支持选择已接入的固定复现测试、明确源码版本和隔离环境，并记录运行条件与结果。当前界面列出的 AUTH-01 测试针对另一组件；本次 libpng 任务未创建验证案例，PoC 状态为 not_run。后续可结合目标构建和输入路径信息补充产品级适用性分析。")

heading("6.8 报告导出与业务闭环")
para("报告中心支持导出 JSON、Markdown、manifest 及复核包。报告关联扫描版本、规则或检测器信息、候选结果、覆盖范围和证据哈希，并纳入已有的人工复核与模型记录。用户可通过提交 SHA 和报告 SHA-256 将导出文件关联到任务，供组件维护者开展后续分析。修复后可对固定版本重新扫描或复验，保留前后端点及复核记录，形成“扫描—复核—验证—报告”的闭环流程。")
para("本次 libpng 任务导出 report.json、report.md、manifest.json、review.json 和 review.md；报告 SHA-256 为 07be8bed379e4df1fe1f7fc17e386ce66305a24003e7dbe8a2559a185d5a79a6。交付内容包含候选及源码证据、解析覆盖、复核状态和后续核验事项，便于复现与持续跟踪。")

para("参考资料", style="Heading 2")
p=para("[1] pnggroup/libpng. Out-of-bounds read/write in the palette expansion on ARM Neon. GitHub Security Advisory GHSA-wjr5-c57x-95m2, 2026-03-25. https://github.com/pnggroup/libpng/security/advisories/GHSA-wjr5-c57x-95m2（访问日期：2026-10-22）。")
p.paragraph_format.first_line_indent=Inches(0)

doc.core_properties.title = "第六章 系统功能实现与用户交互"
doc.core_properties.subject = "中国研究生网络安全创新大赛作品报告独立章节"
doc.core_properties.author = ""
doc.save(OUT)
print(OUT)
